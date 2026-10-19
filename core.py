"""Core functions for a local, deterministic Argument Atlas POC.

No model calls, no firm systems, and no web requests in the default search path.
The initial records are manually curated summaries of *public court orders*.
"""
from __future__ import annotations

import io
import json
import re
import sqlite3
from pathlib import Path
from urllib.parse import urlparse

DATA_PATH = Path(__file__).resolve().parent / "data" / "records.json"
MAX_PDF_BYTES = 20 * 1024 * 1024
ALLOWED_PDF_HOSTS = {"cases.justia.com", "www.govinfo.gov"}


def load_records(path: Path = DATA_PATH) -> list[dict]:
    records = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(records, list) or not records:
        raise ValueError("Expected a nonempty list of records")
    ids = set()
    for item in records:
        for name in ("id", "case_name", "source_url", "source_page", "court_reasoning", "outcome"):
            if not item.get(name):
                raise ValueError(f"Missing field {name!r} in record: {item}")
        if item["id"] in ids:
            raise ValueError(f"Duplicate record ID: {item['id']}")
        ids.add(item["id"])
    return records


def tokenize_query(query: str) -> list[str]:
    # Only actual words become FTS tokens: no user-supplied FTS operator syntax.
    return re.findall(r"[a-zA-Z0-9]{2,}", query.lower())[:12]


def search(records: list[dict], query: str = "", court: str = "All", outcome: str = "All", limit: int = 100) -> list[dict]:
    """SQLite FTS5 ranks curated summaries, with a plain-text fallback."""
    scoped = [r for r in records if (court == "All" or r["court"] == court)
              and (outcome == "All" or r["outcome"] == outcome)]
    terms = tokenize_query(query)
    if not terms:
        return sorted(scoped, key=lambda r: (r["date"], r["id"]), reverse=True)[:limit]

    by_id = {r["id"]: r for r in scoped}
    try:
        db = sqlite3.connect(":memory:")
        try:
            db.execute("CREATE VIRTUAL TABLE knowledge USING fts5(record_id UNINDEXED, body)")
            db.executemany("INSERT INTO knowledge (record_id, body) VALUES (?, ?)", [
                (r["id"], " ".join(str(r.get(k, "")) for k in (
                    "case_name", "court", "expert", "issue", "argument", "opposition", "court_reasoning", "outcome"))
                 + " " + " ".join(r.get("topics", []))) for r in scoped
            ])
            # OR gives broad candidate discovery. It is not semantic similarity.
            fts_query = " OR ".join('"' + t + '"' for t in terms)
            rows = db.execute("SELECT record_id FROM knowledge WHERE knowledge MATCH ? ORDER BY bm25(knowledge) LIMIT ?",
                              (fts_query, limit)).fetchall()
            return [by_id[id_] for (id_,) in rows if id_ in by_id]
        finally:
            db.close()
    except sqlite3.OperationalError as exc:
        if "fts5" not in str(exc).lower():
            raise
        scored = []
        for r in scoped:
            body = " ".join(str(v) for v in r.values()).lower()
            score = sum(body.count(t) for t in terms)
            if score:
                scored.append((score, r))
        return [r for _, r in sorted(scored, key=lambda x: x[0], reverse=True)[:limit]]


def source_link(record: dict) -> str:
    return f"{record['source_url']}#page={int(record['source_page'])}"


def compare_rows(records: list[dict], selected_ids: list[str]) -> list[dict]:
    by_id = {r["id"]: r for r in records}
    return [by_id[id_] for id_ in selected_ids if id_ in by_id]


def download_public_pdf(url: str, cache_dir: Path) -> Path:
    """Optional, explicit public-file retrieval; never accepts arbitrary hosts."""
    from hashlib import sha256
    import requests
    parsed = urlparse(url)
    if parsed.scheme != "https" or parsed.hostname not in ALLOWED_PDF_HOSTS:
        raise ValueError("Only allowlisted public HTTPS PDF hosts are supported")
    cache_dir.mkdir(parents=True, exist_ok=True)
    path = cache_dir / f"{sha256(url.encode()).hexdigest()[:20]}.pdf"
    if path.is_file() and path.stat().st_size:
        return path
    # Avoid redirects to third-party URLs and stream with a size cap.
    with requests.get(url, stream=True, timeout=(8, 35), allow_redirects=False) as response:
        response.raise_for_status()
        with path.with_suffix(".part").open("wb") as handle:
            size = 0
            for chunk in response.iter_content(chunk_size=65536):
                size += len(chunk)
                if size > MAX_PDF_BYTES:
                    handle.close()
                    path.with_suffix(".part").unlink(missing_ok=True)
                    raise ValueError("PDF exceeded maximum allowed size")
                handle.write(chunk)
    temp = path.with_suffix(".part")
    if not temp.read_bytes().startswith(b"%PDF-"):
        temp.unlink(missing_ok=True)
        raise ValueError("The response was not a PDF")
    temp.replace(path)
    return path


def validate_anchor(record: dict, pdf_path: Path) -> tuple[bool, str]:
    """Check the location of a short source phrase, *not* legal claim support."""
    import fitz
    with fitz.open(pdf_path) as doc:
        page_index = int(record["source_page"]) - 1
        if page_index < 0 or page_index >= len(doc):
            return False, "Cited page does not exist in downloaded PDF"
        page_text = doc[page_index].get_text("text")
        def normalize(s):
            return " ".join(re.findall(r"[\w]+", s.lower()))
        if normalize(record["anchor"]) in normalize(page_text):
            return True, "Anchor phrase located on cited PDF page (substantive accuracy not validated)"
        return False, "Anchor phrase not found; inspect the original PDF and check the source page"


def word_comparison(rows: list[dict]) -> bytes:
    """Generate a demonstrative attorney-reviewable *draft*, not legal advice."""
    from docx import Document
    d = Document()
    d.add_heading("Argument Atlas | Expert Challenge Comparison", 0)
    d.add_paragraph("PUBLIC CASE DATA · PROTOTYPE · NOT LEGAL ADVICE · REQUIRES ATTORNEY VERIFICATION")
    d.add_paragraph("Party arguments below are summarized from court orders, not extracted from original briefs. "
                    "Court orders may describe only part of the arguments. Cases are not a statistical sample.")
    for i, row in enumerate(rows):
        if i:
            d.add_page_break()
        d.add_heading(row["case_name"], level=1)
        d.add_paragraph(f"{row['court']} | {row['docket']} | {row['date']} | Expert: {row['expert']}")
        d.add_paragraph(f"Outcome of challenged motion: {row['outcome']} | Motion docket entry: {row['motion_document']}")
        for label, key in (("Issue", "issue"), ("Moving-party argument as described by court", "argument"),
                           ("Opposition as described by court", "opposition"), ("Court's reasoning", "court_reasoning"),
                           ("Scope / limitation", "scope_note")):
            p = d.add_paragraph()
            p.add_run(label + ": ").bold = True
            p.add_run(row[key])
        d.add_paragraph(f"Public court order, document {row['order_document']}, PDF p. {row['source_page']}: {source_link(row)}")
    d.add_heading("Limitations", 1)
    d.add_paragraph("This prototype does not confirm current law, subsequent history, the completeness of case filings, "
                    "or whether a prior decision is persuasive for another matter. Check all sources independently.")
    result = io.BytesIO()
    d.save(result)
    return result.getvalue()
